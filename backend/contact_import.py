"""Bounded spreadsheet parsing. Import never calls Telegram or sends messages."""
import base64,csv,io,re,zipfile
from urllib.parse import urlparse

MAX_BYTES=5*1024*1024
MAX_ROWS=1000
MAX_COLS=30

def cell(value):
    return str(value if value is not None else '').strip()[:500]

def bounded_matrix(rows):
    result=[]
    for row in rows:
        values=[cell(v) for v in row]
        while values and not values[-1]:values.pop()
        if not any(values):continue
        if len(values)>MAX_COLS:raise ValueError('Use at most 30 columns.')
        result.append(values)
        if len(result)>MAX_ROWS+1:raise ValueError('Import at most 1,000 rows at a time.')
    if not result:raise ValueError('The sheet is empty.')
    width=max(map(len,result))
    return [r+['']*(width-len(r)) for r in result]

def read_file(filename,encoded):
    try:data=base64.b64decode(encoded,validate=True)
    except Exception:raise ValueError('The uploaded file could not be read.')
    if not data or len(data)>MAX_BYTES:raise ValueError('Choose a file up to 5 MB.')
    suffix=filename.lower().rsplit('.',1)[-1]
    if suffix in ('csv','tsv'):
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError:raise ValueError('Save the CSV using UTF-8 encoding.')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
        except csv.Error:dialect=csv.excel_tab if suffix=='tsv' else csv.excel
        return bounded_matrix(csv.reader(io.StringIO(text),dialect))
    if suffix!='xlsx':raise ValueError('Choose a CSV, TSV or XLSX file.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist())>500 or sum(x.file_size for x in archive.infolist())>30*1024*1024:raise ValueError('This workbook is too large when expanded. Export the required tab as CSV.')
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            sheet=book.active
            if sheet.max_column>MAX_COLS or sheet.max_row>MAX_ROWS+1:raise ValueError('Use at most 1,000 rows and 30 columns. Export only the required range.')
            return bounded_matrix(tuple('' if c.data_type=='f' else c.value for c in row) for row in sheet.iter_rows())
        finally:book.close()
    except ValueError:raise
    except Exception:raise ValueError('Could not open this XLSX workbook. Try exporting it as CSV.')

ALIASES={
 'telegram':{'telegram','telegramlink','telegramusername','telegramhandle','username','handle','profile','link','telegramurl','tg','tme','usernames','telegramlinks'},
 'name':{'name','fullname','contactname','projectname'},
 'role':{'role','title','roletitle','position'},
 'group':{'group','community','project','company'},
 'kind':{'type','kind','contacttype'},
 'country':{'country','countrycode','phonecountry'},
 'timezone':{'timezone','ianatimezone','tz'},
}
def header_key(value):return re.sub('[^a-z0-9]','',value.casefold())
def default_mapping(rows):
    first=rows[0];mapping={}
    for field,aliases in ALIASES.items():
        matches=[i for i,h in enumerate(first) if header_key(h) in aliases]
        mapping[field]=matches[0] if matches else None
    has_header=mapping['telegram'] is not None
    if not has_header:mapping['telegram']=0
    return mapping,has_header

def telegram_handle(value):
    value=cell(value).lstrip("'").strip()
    if re.match(r'^(?:https?://)?(?:www\.)?(?:t\.me|telegram\.me)/',value,re.I):
        u=urlparse(value if '://' in value else 'https://'+value)
        if u.username or u.password or u.query or u.fragment:raise ValueError('Use a public Telegram profile or group link without extra parameters.')
        value=u.path.strip('/')
    value=value.lstrip('@')
    if value.lower().endswith('bot'):raise ValueError('Bot usernames are excluded.')
    if not re.fullmatch('[A-Za-z][A-Za-z0-9_]{3,31}',value):raise ValueError('Use @username or https://t.me/username. Private invites and message links must be scanned in Find team.')
    return value.lower()

def sheet_identifier(url):
    u=urlparse(url)
    match=re.fullmatch(r'/spreadsheets/d/([A-Za-z0-9_-]+)(?:/.*)?',u.path)
    if u.scheme!='https' or u.hostname!='docs.google.com' or u.username or u.password or u.port or not match:raise ValueError('Paste a Google Sheets URL beginning https://docs.google.com/spreadsheets/d/.')
    return match[1]
